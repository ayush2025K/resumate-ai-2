"""
Document extraction service for Resumate AI.
Supports PDF (via PyMuPDF) and DOCX (via python-docx) with robust validation,
magic byte verification, error handling, and intelligent section segmentation.
"""
import io
import re
import pymupdf as fitz
import docx
from typing import Dict, Tuple, Optional, Any

# Maximum file size: 10 MB
MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024

# Magic byte signatures
PDF_MAGIC_BYTES = b"%PDF"
DOCX_MAGIC_BYTES = b"PK\x03\x04"

# Section Header Regular Expressions
SECTION_PATTERNS = {
    "summary": re.compile(
        r"^(?:professional\s+summary|executive\s+summary|career\s+summary|about\s+me|profile|objective|summary)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "skills": re.compile(
        r"^(?:technical\s+skills|core\s+competencies|technologies|skills\s+(?:&|and)\s+abilities|toolkit|areas\s+of\s+expertise|skills)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "experience": re.compile(
        r"^(?:work\s+experience|professional\s+experience|employment\s+history|experience|internships|work\s+history)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "education": re.compile(
        r"^(?:academic\s+background|education\s+(?:&|and)\s+qualifications|educational\s+background|academic\s+history|education|qualifications)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "projects": re.compile(
        r"^(?:key\s+projects|academic\s+projects|personal\s+projects|selected\s+projects|technical\s+projects|projects)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "certifications": re.compile(
        r"^(?:certifications\s+(?:&|and)\s+licenses|licenses\s+(?:&|and)\s+certifications|courses\s+(?:&|and)\s+certifications|certificates|certifications)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
    "achievements": re.compile(
        r"^(?:honors\s+(?:&|and)\s+awards|awards\s+(?:&|and)\s+achievements|extracurricular\s+activities|achievements|publications)[\s:]*$",
        re.IGNORECASE | re.MULTILINE
    ),
}


class DocumentExtractionError(Exception):
    """Custom exception for clean error reporting without exposing internals."""
    pass


def validate_file_content(filename: str, content: bytes) -> str:
    """
    Validates file size, extension, and magic byte signature.
    Returns detected file type ('pdf' or 'docx').
    """
    if not content or len(content) == 0:
        raise DocumentExtractionError(f"File '{filename}' is empty (0 bytes). Please upload a valid document.")

    if len(content) > MAX_FILE_SIZE_BYTES:
        max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
        size_mb = round(len(content) / (1024 * 1024), 2)
        raise DocumentExtractionError(
            f"File '{filename}' exceeds maximum allowed size of {max_mb}MB (Received: {size_mb}MB)."
        )

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        if not content.startswith(PDF_MAGIC_BYTES):
            raise DocumentExtractionError(
                f"File '{filename}' has a .pdf extension but lacks a valid PDF file signature."
            )
        return "pdf"
    elif lower_name.endswith(".docx"):
        if not content.startswith(DOCX_MAGIC_BYTES):
            raise DocumentExtractionError(
                f"File '{filename}' has a .docx extension but lacks a valid DOCX file signature."
            )
        return "docx"
    else:
        raise DocumentExtractionError(
            f"Unsupported file format for '{filename}'. Only .pdf and .docx files are accepted."
        )


def clean_extracted_text(text: str) -> str:
    """
    Standardizes whitespace, line breaks, bullet markers, and unicode symbols.
    """
    if not text:
        return ""
    # Normalize bullet points to a standard dash
    text = re.sub(r"[\u2022\u2023\u25E6\u2043\u2219\u25CF\u25AA\u25AB\u25B6\u27A4]", "\n- ", text)
    # Replace non-breaking spaces and zero-width spaces
    text = text.replace("\u00a0", " ").replace("\u200b", "").replace("\ufeff", "")
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse 3 or more consecutive newlines to 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(content: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text and metadata from PDF bytes using PyMuPDF (fitz).
    """
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to open PDF document. The file may be corrupt or encrypted.") from None

    try:
        page_texts = []
        page_count = len(doc)
        if page_count == 0:
            raise DocumentExtractionError("PDF file contains no pages.")

        for page_num in range(page_count):
            page = doc.load_page(page_num)
            page_text = page.get_text("text")
            if page_text:
                page_texts.append(page_text)

        full_text = "\n".join(page_texts)
        cleaned_text = clean_extracted_text(full_text)

        if not cleaned_text or len(cleaned_text.strip()) < 10:
            raise DocumentExtractionError(
                "Could not extract readable text from PDF. The document may be scanned, image-only, or empty."
            )

        metadata = {
            "page_count": page_count,
            "char_count": len(cleaned_text),
            "word_count": len(cleaned_text.split()),
            "format": "pdf",
            "pdf_title": doc.metadata.get("title") if doc.metadata else None,
            "pdf_author": doc.metadata.get("author") if doc.metadata else None,
        }
        return cleaned_text, metadata
    finally:
        doc.close()


def extract_text_from_docx(content: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts text and table contents from DOCX bytes using python-docx.
    """
    try:
        doc_stream = io.BytesIO(content)
        doc = docx.Document(doc_stream)
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to parse DOCX document. File may be damaged or password protected.") from None

    parts = []
    # Extract paragraphs
    for p in doc.paragraphs:
        p_text = p.text.strip()
        if p_text:
            parts.append(p_text)

    # Extract tables (many resumes format experience or skills in tables)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Deduplicate identical adjacent cells (due to cell merging)
                deduped = []
                for cell in row_cells:
                    if not deduped or cell != deduped[-1]:
                        deduped.append(cell)
                parts.append(" | ".join(deduped))

    full_text = "\n".join(parts)
    cleaned_text = clean_extracted_text(full_text)

    if not cleaned_text or len(cleaned_text.strip()) < 10:
        raise DocumentExtractionError("DOCX file contains no readable text content.")

    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "char_count": len(cleaned_text),
        "word_count": len(cleaned_text.split()),
        "format": "docx",
    }
    return cleaned_text, metadata


def extract_document(filename: str, content: bytes) -> Tuple[str, Dict[str, Any], Dict[str, str]]:
    """
    Validates and extracts raw text, metadata, and segmented sections from a resume file.
    Returns: (raw_text, metadata, sections_dict)
    """
    file_type = validate_file_content(filename, content)
    if file_type == "pdf":
        raw_text, metadata = extract_text_from_pdf(content)
    else:
        raw_text, metadata = extract_text_from_docx(content)

    sections = segment_sections(raw_text)
    return raw_text, metadata, sections


def segment_sections(text: str) -> Dict[str, str]:
    """
    Intelligently splits resume text into logical sections based on detected headers.
    """
    lines = text.split("\n")
    header_indices = []

    for idx, line in enumerate(lines):
        clean_line = line.strip()
        # Header candidates are typically short (under 50 chars) and match section keywords
        if clean_line and len(clean_line) <= 50:
            for sec_name, pattern in SECTION_PATTERNS.items():
                if pattern.match(clean_line):
                    header_indices.append((idx, sec_name, clean_line))
                    break

    sections: Dict[str, str] = {}
    if not header_indices:
        # If no explicit headers were detected, treat the entire document as general content
        sections["general"] = text
        return sections

    # First section before any header (typically contact info & name)
    first_idx = header_indices[0][0]
    if first_idx > 0:
        header_text = "\n".join(lines[:first_idx]).strip()
        if header_text:
            sections["header"] = header_text

    # Extract text between consecutive headers
    for i in range(len(header_indices)):
        current_idx, sec_name, _ = header_indices[i]
        next_idx = header_indices[i + 1][0] if i + 1 < len(header_indices) else len(lines)
        section_body = "\n".join(lines[current_idx + 1 : next_idx]).strip()
        # If section already exists (e.g. multiple project sections), append it
        if sec_name in sections:
            sections[sec_name] += "\n\n" + section_body
        else:
            sections[sec_name] = section_body

    return sections
